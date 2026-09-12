"""El libro Excel como almacén: creación, carga, escritura atómica y copias.

Reglas de juego de docs/03: la interfaz nunca toca openpyxl; el libro se
mantiene cerrado entre operaciones; antes de cada escritura se hace copia de
seguridad; y si el fichero cambió por fuera desde la última lectura, se avisa
en lugar de sobreescribir.
"""

from __future__ import annotations

import os
import shutil
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from ..nucleo import log
from .esquema import TABLAS, Columna, Tabla, Tipo

_log = log.obtener("datos.libro")

AZUL = "4A7FA5"
AZUL_CLARO = "D9E4EC"
GRIS = "F2F2F2"


class LibroBloqueado(Exception):
    """El libro está abierto en Excel y no se puede escribir."""


class LibroModificadoFuera(Exception):
    """El libro cambió en disco desde que lo cargamos."""


class EsquemaDesactualizado(Exception):
    """Al libro le faltan hojas o columnas que el programa sí espera.

    Pasa cuando el libro se creó con una versión anterior. No es un error del
    usuario ni hay que rehacer nada: se migra añadiendo lo que falta.
    """

    def __init__(self, faltan: dict[str, list[str]]):
        self.faltan = faltan
        partes = [f"{hoja}: {', '.join(columnas) if columnas else 'la hoja entera'}"
                  for hoja, columnas in faltan.items()]
        super().__init__("Al libro le falta — " + " · ".join(partes))

    def resumen(self) -> str:
        hojas = [h for h, c in self.faltan.items() if not c]
        columnas = sum(len(c) for c in self.faltan.values())
        partes = []
        if hojas:
            partes.append(f"{len(hojas)} hojas")
        if columnas:
            partes.append(f"{columnas} columnas")
        return " y ".join(partes) or "nada"


# --- Conversión de valores ---------------------------------------------------

def a_celda(valor, columna: Columna):
    """Valor de Python -> valor que se escribe en la celda."""
    if valor is None or valor == "":
        return None
    if columna.tipo is Tipo.BOOL:
        if isinstance(valor, str):
            return "SI" if valor.strip().upper() in ("SI", "SÍ", "S", "1", "TRUE") else "NO"
        return "SI" if valor else "NO"
    if columna.tipo is Tipo.FECHA:
        return valor if isinstance(valor, (date, datetime)) else str(valor)
    if columna.tipo is Tipo.FECHA_HORA:
        return valor if isinstance(valor, datetime) else str(valor)
    if columna.tipo is Tipo.ENTERO:
        return int(valor)
    if columna.tipo is Tipo.DECIMAL:
        return float(valor)
    return str(valor)


def de_celda(valor, columna: Columna):
    """Valor de la celda -> valor de Python."""
    if valor is None or (isinstance(valor, str) and not valor.strip()):
        return None
    if columna.tipo is Tipo.BOOL:
        return str(valor).strip().upper() in ("SI", "SÍ", "S", "1", "TRUE", "VERDADERO")
    if columna.tipo is Tipo.FECHA:
        return valor.date() if isinstance(valor, datetime) else valor
    if columna.tipo is Tipo.ENTERO:
        return int(valor)
    if columna.tipo is Tipo.DECIMAL:
        return float(valor)
    if columna.tipo is Tipo.FECHA_HORA:
        return valor
    return str(valor).strip()


# --- Creación del libro ------------------------------------------------------

def _formatear_hoja(hoja, tabla: Tabla) -> None:
    cabecera_fondo = PatternFill("solid", fgColor=AZUL)
    cabecera_letra = Font(bold=True, color="FFFFFF", size=10)
    borde = Border(bottom=Side(style="thin", color=AZUL))

    for i, columna in enumerate(tabla.columnas, start=1):
        celda = hoja.cell(row=1, column=i, value=columna.nombre)
        celda.fill = cabecera_fondo
        celda.font = cabecera_letra
        celda.border = borde
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        if columna.ayuda:
            celda.comment = None  # la ayuda va en la hoja _ESQUEMA, más visible
        hoja.column_dimensions[get_column_letter(i)].width = columna.ancho

        if columna.tipo is Tipo.LISTA:
            opciones = ",".join(columna.opciones)
            if len(opciones) <= 250:  # límite de las listas en línea de Excel
                dv = DataValidation(type="list", formula1=f'"{opciones}"', allow_blank=True)
                hoja.add_data_validation(dv)
                dv.add(f"{get_column_letter(i)}2:{get_column_letter(i)}5000")
        elif columna.tipo is Tipo.BOOL:
            dv = DataValidation(type="list", formula1='"SI,NO"', allow_blank=True)
            hoja.add_data_validation(dv)
            dv.add(f"{get_column_letter(i)}2:{get_column_letter(i)}5000")

    hoja.row_dimensions[1].height = 30
    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = (
        f"A1:{get_column_letter(len(tabla.columnas))}1"
    )


def _hoja_esquema(libro: Workbook) -> None:
    """Hoja de referencia legible: qué es cada tabla y cada columna."""
    hoja = libro.create_sheet("_ESQUEMA", 0)
    hoja.sheet_properties.tabColor = AZUL
    cabeceras = ["Grupo", "Tabla", "Columna", "Tipo", "Obligatorio", "Opciones / Referencia", "Ayuda"]
    anchos = [16, 24, 30, 14, 12, 40, 45]
    for i, (texto, ancho) in enumerate(zip(cabeceras, anchos), start=1):
        celda = hoja.cell(row=1, column=i, value=texto)
        celda.fill = PatternFill("solid", fgColor=AZUL)
        celda.font = Font(bold=True, color="FFFFFF")
        hoja.column_dimensions[get_column_letter(i)].width = ancho

    fila = 2
    for tabla in TABLAS:
        for columna in tabla.columnas:
            detalle = ""
            if columna.opciones:
                detalle = " | ".join(columna.opciones)
            elif columna.ref:
                detalle = f"-> {columna.ref}"
            hoja.cell(row=fila, column=1, value=tabla.grupo)
            hoja.cell(row=fila, column=2, value=tabla.nombre)
            hoja.cell(row=fila, column=3, value=columna.nombre)
            hoja.cell(row=fila, column=4, value=columna.tipo.value)
            hoja.cell(row=fila, column=5, value="SI" if columna.obligatorio else "")
            hoja.cell(row=fila, column=6, value=detalle)
            hoja.cell(row=fila, column=7, value=columna.ayuda)
            if fila % 2 == 0:
                for col in range(1, 8):
                    hoja.cell(row=fila, column=col).fill = PatternFill("solid", fgColor=GRIS)
            fila += 1
    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = f"A1:G{fila - 1}"


def crear(ruta: Path, sobreescribir: bool = False) -> Path:
    """Genera el libro con todas las hojas, cabeceras y validaciones."""
    ruta = Path(ruta)
    if ruta.exists() and not sobreescribir:
        raise FileExistsError(f"Ya existe {ruta}. Usa sobreescribir=True si es lo que quieres.")
    libro = Workbook()
    libro.remove(libro.active)
    for tabla in TABLAS:
        hoja = libro.create_sheet(tabla.nombre)
        _formatear_hoja(hoja, tabla)
    _hoja_esquema(libro)
    ruta.parent.mkdir(parents=True, exist_ok=True)
    libro.save(ruta)
    _log.info("Libro creado en %s con %d hojas", ruta, len(TABLAS))
    return ruta


# --- Carga y guardado --------------------------------------------------------

def _fichero_bloqueo(ruta: Path) -> Path:
    return ruta.parent / f"~${ruta.name}"


def esta_abierto_en_excel(ruta: Path) -> bool:
    return _fichero_bloqueo(ruta).exists()


class Libro:
    """Carga el libro en memoria, y lo escribe de vuelta de forma atómica."""

    def __init__(self, ruta: Path, carpeta_backups: Path, copias_a_conservar: int = 30):
        self.ruta = Path(ruta)
        self.carpeta_backups = Path(carpeta_backups)
        self.copias_a_conservar = copias_a_conservar
        self.filas: dict[str, list[dict]] = {}
        self._mtime: float | None = None

    # -- lectura --
    def cargar(self) -> None:
        if not self.ruta.exists():
            raise FileNotFoundError(
                f"No existe el libro {self.ruta}. Créalo con herramientas/crear_libro.py"
            )
        libro = load_workbook(self.ruta, data_only=True)
        pendiente = _que_falta(libro)
        if pendiente:
            libro.close()
            raise EsquemaDesactualizado(pendiente)

        self.filas = {}
        for tabla in TABLAS:
            hoja = libro[tabla.nombre]
            cabeceras = [c.value for c in hoja[1]]
            indice = {nombre: i for i, nombre in enumerate(cabeceras)}
            registros: list[dict] = []
            for fila in hoja.iter_rows(min_row=2, values_only=True):
                if all(v is None or v == "" for v in fila):
                    continue
                registro = {}
                for columna in tabla.columnas:
                    pos = indice[columna.nombre]
                    valor = fila[pos] if pos < len(fila) else None
                    registro[columna.nombre] = de_celda(valor, columna)
                registros.append(registro)
            self.filas[tabla.nombre] = registros
        libro.close()
        self._mtime = self.ruta.stat().st_mtime
        total = sum(len(v) for v in self.filas.values())
        _log.info("Libro cargado: %d registros en %d hojas", total, len(self.filas))

    # -- escritura --
    def _copia_seguridad(self) -> Path | None:
        if not self.ruta.exists():
            return None
        self.carpeta_backups.mkdir(parents=True, exist_ok=True)
        marca = datetime.now().strftime("%Y%m%d_%H%M%S")
        destino = self.carpeta_backups / f"{self.ruta.stem}_{marca}{self.ruta.suffix}"
        shutil.copy2(self.ruta, destino)
        self._podar_copias()
        return destino

    def _podar_copias(self) -> None:
        copias = sorted(
            self.carpeta_backups.glob(f"{self.ruta.stem}_*{self.ruta.suffix}"),
            key=lambda p: p.name,
        )
        for vieja in copias[: max(0, len(copias) - self.copias_a_conservar)]:
            vieja.unlink(missing_ok=True)

    def guardar(self, forzar: bool = False) -> None:
        if esta_abierto_en_excel(self.ruta):
            raise LibroBloqueado(
                "El libro está abierto en Excel. Ciérralo y vuelve a intentarlo."
            )
        if not forzar and self._mtime is not None and self.ruta.exists():
            if self.ruta.stat().st_mtime > self._mtime + 1:
                raise LibroModificadoFuera(
                    "El libro ha cambiado en disco desde que se cargó. "
                    "Si guardas ahora perderás esos cambios."
                )
        self._copia_seguridad()

        libro = Workbook()
        libro.remove(libro.active)
        for tabla in TABLAS:
            hoja = libro.create_sheet(tabla.nombre)
            _formatear_hoja(hoja, tabla)
            for n, registro in enumerate(self.filas.get(tabla.nombre, []), start=2):
                for i, columna in enumerate(tabla.columnas, start=1):
                    hoja.cell(row=n, column=i,
                              value=a_celda(registro.get(columna.nombre), columna))
        _hoja_esquema(libro)

        temporal = self.ruta.with_suffix(self.ruta.suffix + ".tmp")
        libro.save(temporal)
        os.replace(temporal, self.ruta)     # atómico dentro del mismo volumen
        self._mtime = self.ruta.stat().st_mtime
        _log.info("Libro guardado en %s", self.ruta)


# --- Migración del esquema ---------------------------------------------------

def _que_falta(libro: Workbook) -> dict[str, list[str]]:
    """Hojas y columnas que el esquema espera y el libro no tiene.

    Una hoja ausente se devuelve con la lista vacía; una hoja presente, con las
    columnas que le falten. Las columnas de más no se tocan: pueden ser tuyas.
    """
    faltan: dict[str, list[str]] = {}
    for tabla in TABLAS:
        if tabla.nombre not in libro.sheetnames:
            faltan[tabla.nombre] = []
            continue
        cabeceras = [c.value for c in libro[tabla.nombre][1]]
        ausentes = [c for c in tabla.nombres_columnas if c not in cabeceras]
        if ausentes:
            faltan[tabla.nombre] = ausentes
    return faltan


def revisar(ruta: Path) -> dict[str, list[str]]:
    libro = load_workbook(Path(ruta), read_only=False)
    try:
        return _que_falta(libro)
    finally:
        libro.close()


def migrar(ruta: Path, carpeta_backups: Path | None = None) -> list[str]:
    """Añade al libro las hojas y columnas que falten, sin tocar los datos.

    Las columnas nuevas se añaden al final de la cabecera, que es donde no
    molestan: la aplicación lee por nombre, no por posición.
    """
    ruta = Path(ruta)
    if esta_abierto_en_excel(ruta):
        raise LibroBloqueado(
            "El libro está abierto en Excel. Ciérralo para poder actualizarlo.")

    libro = load_workbook(ruta)
    pendiente = _que_falta(libro)
    if not pendiente:
        libro.close()
        return []

    if carpeta_backups:
        carpeta_backups = Path(carpeta_backups)
        carpeta_backups.mkdir(parents=True, exist_ok=True)
        marca = datetime.now().strftime("%Y%m%d_%H%M%S")
        shutil.copy2(ruta, carpeta_backups / f"{ruta.stem}_previo_{marca}{ruta.suffix}")

    hecho: list[str] = []
    for nombre, columnas in pendiente.items():
        tabla = next(t for t in TABLAS if t.nombre == nombre)
        if not columnas:
            _formatear_hoja(libro.create_sheet(nombre), tabla)
            hecho.append(f"hoja {nombre} creada")
            continue
        hoja = libro[nombre]
        cabeceras = [c.value for c in hoja[1]]
        for columna in columnas:
            definicion = tabla.columna(columna)
            posicion = len(cabeceras) + 1
            celda = hoja.cell(row=1, column=posicion, value=columna)
            celda.fill = PatternFill("solid", fgColor=AZUL)
            celda.font = Font(bold=True, color="FFFFFF", size=10)
            celda.alignment = Alignment(horizontal="center", vertical="center",
                                        wrap_text=True)
            hoja.column_dimensions[get_column_letter(posicion)].width = definicion.ancho
            if definicion.tipo is Tipo.BOOL:
                dv = DataValidation(type="list", formula1='"SI,NO"', allow_blank=True)
                hoja.add_data_validation(dv)
                dv.add(f"{get_column_letter(posicion)}2:{get_column_letter(posicion)}5000")
            elif definicion.tipo is Tipo.LISTA:
                opciones = ",".join(definicion.opciones)
                if len(opciones) <= 250:
                    dv = DataValidation(type="list", formula1=f'"{opciones}"',
                                        allow_blank=True)
                    hoja.add_data_validation(dv)
                    dv.add(f"{get_column_letter(posicion)}2:"
                           f"{get_column_letter(posicion)}5000")
            cabeceras.append(columna)
            hecho.append(f"{nombre}.{columna}")

    if "_ESQUEMA" in libro.sheetnames:
        del libro["_ESQUEMA"]
    _hoja_esquema(libro)

    temporal = ruta.with_suffix(ruta.suffix + ".tmp")
    libro.save(temporal)
    os.replace(temporal, ruta)
    _log.info("Libro migrado: %s", ", ".join(hecho))
    return hecho


def abrir(ruta: Path, carpeta_backups: Path, copias_a_conservar: int = 30,
          migrar_si_hace_falta: bool = True) -> Libro:
    """Carga el libro, actualizando antes su esquema si se quedó atrás."""
    libro = Libro(ruta, carpeta_backups, copias_a_conservar)
    try:
        libro.cargar()
    except EsquemaDesactualizado:
        if not migrar_si_hace_falta:
            raise
        cambios = migrar(ruta, carpeta_backups)
        _log.info("Esquema actualizado antes de abrir: %d cambios", len(cambios))
        libro.cargar()
    return libro
