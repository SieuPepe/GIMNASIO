"""Pruebas de la capa de datos y de la lógica de usuarios.

    python -m unittest discover -s pruebas
"""

import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from zeu.datos import libro as ml, semillas                       # noqa: E402
from zeu.datos.esquema import TABLAS, Tipo, tabla                  # noqa: E402
from zeu.datos.repositorio import ErrorValidacion, Repositorio     # noqa: E402
from zeu.nucleo.fechas import a_fecha, edad                        # noqa: E402
from zeu.servicios import usuarios as srv                          # noqa: E402


class Esquema(unittest.TestCase):
    def test_referencias_apuntan_a_tablas_reales(self):
        nombres = {t.nombre for t in TABLAS}
        for t in TABLAS:
            for c in t.columnas:
                if c.tipo is Tipo.REF:
                    self.assertIn(c.ref, nombres, f"{t.nombre}.{c.nombre}")

    def test_prefijos_unicos(self):
        prefijos = [t.prefijo for t in TABLAS if t.prefijo]
        self.assertEqual(len(prefijos), len(set(prefijos)))

    def test_la_clave_es_la_primera_columna(self):
        self.assertEqual(tabla("T_USUARIOS").clave, "ID_Usuario")


class Base(unittest.TestCase):
    def setUp(self):
        self.temporal = tempfile.TemporaryDirectory()
        carpeta = Path(self.temporal.name)
        self.ruta = carpeta / "datos.xlsx"
        self.backups = carpeta / "backups"
        ml.crear(self.ruta)
        self.libro = ml.Libro(self.ruta, self.backups, copias_a_conservar=3)
        self.libro.cargar()
        semillas.sembrar(self.libro)
        self.repo = Repositorio(self.libro)

    def tearDown(self):
        self.temporal.cleanup()

    def alta_ana(self) -> str:
        return srv.alta(self.repo, {
            "Nombre": "Ana", "Apellidos": "García", "F_Nacimiento": date(1985, 3, 12),
            "Sexo": "M", "Email": "ana@ejemplo.com"})


class Repositorio_(Base):
    def test_alta_asigna_identificador_correlativo(self):
        self.assertEqual(self.alta_ana(), "USR-0001")
        self.assertEqual(self.alta_ana(), "USR-0002")

    def test_un_alta_rechazada_no_consume_identificador(self):
        with self.assertRaises(ErrorValidacion):
            self.repo.insertar("T_USUARIOS", {"Nombre": "X", "Sexo": "Z"})
        self.assertEqual(self.alta_ana(), "USR-0001")

    def test_valida_obligatorios_y_listas(self):
        with self.assertRaises(ErrorValidacion) as caso:
            self.repo.insertar("T_USUARIOS", {"Nombre": "X", "Sexo": "Z"})
        errores = " ".join(caso.exception.errores)
        self.assertIn("F_Nacimiento", errores)
        self.assertIn("Sexo", errores)

    def test_valida_referencias(self):
        with self.assertRaises(ErrorValidacion) as caso:
            self.repo.insertar("T_OBJETIVOS", {
                "ID_Usuario": "USR-9999", "F_Registro": date.today(),
                "Tipo": "FUERZA", "Descripcion": "x", "Estado": "ACTIVO"})
        self.assertIn("USR-9999", " ".join(caso.exception.errores))

    def test_no_borra_lo_que_esta_en_uso(self):
        uid = self.alta_ana()
        self.repo.insertar("T_OBJETIVOS", {
            "ID_Usuario": uid, "F_Registro": date.today(), "Tipo": "FUERZA",
            "Descripcion": "Mejorar sentadilla", "Estado": "ACTIVO"})
        with self.assertRaises(ErrorValidacion):
            self.repo.borrar("T_USUARIOS", uid)
        self.assertTrue(self.repo.existe("T_USUARIOS", uid))

    def test_config_sembrada(self):
        self.assertEqual(self.repo.config("dias_aviso_previo"), "14")


class Persistencia(Base):
    def test_ida_y_vuelta_conserva_tipos(self):
        uid = self.alta_ana()
        self.repo.insertar("T_SALUD", {
            "ID_Usuario": uid, "F_Registro": date(2026, 9, 1), "PARQ_1_SN": True,
            "PARQ_2_SN": False, "TA_Sistolica": 152, "FC_Reposo": 78})
        self.libro.guardar()

        otro = ml.Libro(self.ruta, self.backups)
        otro.cargar()
        repo = Repositorio(otro)
        salud = repo.listar("T_SALUD")[0]
        self.assertIs(salud["PARQ_1_SN"], True)
        self.assertIs(salud["PARQ_2_SN"], False)
        self.assertEqual(salud["TA_Sistolica"], 152)
        self.assertEqual(salud["F_Registro"], date(2026, 9, 1))
        self.assertEqual(repo.obtener("T_USUARIOS", uid)["F_Nacimiento"], date(1985, 3, 12))

    def test_guardar_hace_copia_y_poda(self):
        for _ in range(5):
            self.alta_ana()
            self.libro.guardar()
        copias = list(self.backups.glob("*.xlsx"))
        self.assertLessEqual(len(copias), 3)
        self.assertGreater(len(copias), 0)

    def test_avisa_si_el_libro_esta_abierto_en_excel(self):
        bloqueo = self.ruta.parent / f"~${self.ruta.name}"
        bloqueo.write_text("")
        with self.assertRaises(ml.LibroBloqueado):
            self.libro.guardar()
        bloqueo.unlink()

    def test_avisa_si_cambio_por_fuera(self):
        self.libro.guardar()
        self.libro._mtime = 0.0        # como si lo hubiéramos leído hace mucho
        with self.assertRaises(ml.LibroModificadoFuera):
            self.libro.guardar()
        self.libro.guardar(forzar=True)


class Migracion(Base):
    """Un libro creado con una versión anterior se actualiza sin perder datos."""

    def _envejecer(self, quitar: dict[str, list[str]]) -> None:
        from openpyxl import load_workbook
        libro = load_workbook(self.ruta)
        for hoja, columnas in quitar.items():
            pagina = libro[hoja]
            for nombre in columnas:
                posicion = [c.value for c in pagina[1]].index(nombre) + 1
                pagina.delete_cols(posicion)
        libro.save(self.ruta)
        libro.close()

    def test_detecta_lo_que_falta(self):
        self.libro.guardar()
        self._envejecer({"T_CAT_TESTS": ["Calculada_SN", "Decimales"]})
        self.assertEqual(ml.revisar(self.ruta),
                         {"T_CAT_TESTS": ["Calculada_SN", "Decimales"]})

    def test_cargar_avisa_en_vez_de_reventar(self):
        self.libro.guardar()
        self._envejecer({"T_COLA_MAIL": ["Quiere_Cita"]})
        with self.assertRaises(ml.EsquemaDesactualizado) as caso:
            ml.Libro(self.ruta, self.backups).cargar()
        self.assertIn("T_COLA_MAIL", caso.exception.faltan)

    def test_migrar_conserva_los_datos(self):
        uid = self.alta_ana()
        self.repo.insertar("T_OBJETIVOS", {
            "ID_Usuario": uid, "F_Registro": date(2026, 5, 4), "Tipo": "FUERZA",
            "Descripcion": "Mejorar la sentadilla", "Estado": "ACTIVO"})
        self.libro.guardar()
        self._envejecer({"T_CAT_TESTS": ["Calculada_SN", "Decimales"],
                         "T_COLA_MAIL": ["Quiere_Cita"]})

        cambios = ml.migrar(self.ruta, self.backups)
        self.assertEqual(len(cambios), 3)

        otro = ml.Libro(self.ruta, self.backups)
        otro.cargar()
        repo = Repositorio(otro)
        usuario = repo.obtener("T_USUARIOS", uid)
        self.assertEqual(usuario["Nombre"], "Ana")
        self.assertEqual(usuario["F_Nacimiento"], date(1985, 3, 12))
        self.assertEqual(repo.listar("T_OBJETIVOS")[0]["Descripcion"],
                         "Mejorar la sentadilla")

    def test_migrar_hace_copia_previa(self):
        self.libro.guardar()
        self._envejecer({"T_COLA_MAIL": ["Quiere_Cita"]})
        ml.migrar(self.ruta, self.backups)
        self.assertTrue(list(self.backups.glob("*previo*.xlsx")))

    def test_migrar_dos_veces_no_hace_nada_la_segunda(self):
        self.libro.guardar()
        self._envejecer({"T_COLA_MAIL": ["Quiere_Cita"]})
        self.assertEqual(len(ml.migrar(self.ruta, self.backups)), 1)
        self.assertEqual(ml.migrar(self.ruta, self.backups), [])

    def test_no_migra_con_el_libro_abierto_en_excel(self):
        self.libro.guardar()
        self._envejecer({"T_COLA_MAIL": ["Quiere_Cita"]})
        bloqueo = self.ruta.parent / f"~${self.ruta.name}"
        bloqueo.write_text("")
        with self.assertRaises(ml.LibroBloqueado):
            ml.migrar(self.ruta, self.backups)
        bloqueo.unlink()

    def test_abrir_migra_por_su_cuenta(self):
        uid = self.alta_ana()
        self.libro.guardar()
        self._envejecer({"T_CAT_TESTS": ["Decimales"]})
        libro = ml.abrir(self.ruta, self.backups)
        self.assertEqual(Repositorio(libro).obtener("T_USUARIOS", uid)["Nombre"], "Ana")

    def test_las_columnas_de_mas_no_se_tocan(self):
        """Si alguien añade una columna suya al Excel, se respeta."""
        from openpyxl import load_workbook
        self.libro.guardar()
        libro = load_workbook(self.ruta)
        hoja = libro["T_USUARIOS"]
        hoja.cell(row=1, column=hoja.max_column + 1, value="Mi_Columna")
        libro.save(self.ruta)
        libro.close()
        self._envejecer({"T_COLA_MAIL": ["Quiere_Cita"]})
        ml.migrar(self.ruta, self.backups)
        libro = load_workbook(self.ruta)
        self.assertIn("Mi_Columna", [c.value for c in libro["T_USUARIOS"][1]])
        libro.close()


class Cribado(unittest.TestCase):
    def test_tension_se_clasifica_por_el_criterio_europeo(self):
        self.assertEqual(srv.clasificar_tension(118, 76)[0], "OPTIMA")
        self.assertEqual(srv.clasificar_tension(134, 82)[0], "NORMAL_ALTA")
        self.assertEqual(srv.clasificar_tension(145, 88)[1], 1)
        self.assertEqual(srv.clasificar_tension(120, 105)[1], 2)
        self.assertEqual(srv.clasificar_tension(None, None)[0], "SIN_DATO")

    def test_parq_critico_exige_informe(self):
        resultado = srv.evaluar_cribado({"PARQ_2_SN": True})
        self.assertTrue(resultado.requiere_informe)
        self.assertTrue(resultado.bloquea_esfuerzo_maximo)
        self.assertFalse(resultado.apto)

    def test_parq_no_critico_bloquea_pero_no_exige_informe(self):
        resultado = srv.evaluar_cribado({"PARQ_6_SN": True})
        self.assertFalse(resultado.requiere_informe)
        self.assertTrue(resultado.bloquea_esfuerzo_maximo)

    def test_sin_hallazgos_es_apto(self):
        self.assertTrue(srv.evaluar_cribado({"TA_Sistolica": 118,
                                             "TA_Diastolica": 74,
                                             "FC_Reposo": 64}).apto)


class Fechas(unittest.TestCase):
    def test_edad_cumplida(self):
        self.assertEqual(edad(date(1985, 3, 12), date(2026, 3, 11)), 40)
        self.assertEqual(edad(date(1985, 3, 12), date(2026, 3, 12)), 41)

    def test_acepta_los_formatos_habituales(self):
        self.assertEqual(a_fecha("12/03/1985"), date(1985, 3, 12))
        self.assertEqual(a_fecha("1985-03-12"), date(1985, 3, 12))
        self.assertIsNone(a_fecha(""))


class Avisos(Base):
    def test_lista_lo_que_falta_de_cada_usuario(self):
        uid = self.alta_ana()
        pendientes = srv.avisos(self.repo, self.repo.obtener("T_USUARIOS", uid))
        self.assertEqual(len(pendientes), 4)   # consentimiento, salud, objetivo, disponibilidad

    def test_ficha_completa_no_deja_avisos(self):
        uid = self.alta_ana()
        self.repo.actualizar("T_USUARIOS", uid, {
            "Consentimiento_SN": True, "F_Consentimiento": date.today(),
            "Origen_Consentimiento": "FORM"})
        self.repo.insertar("T_SALUD", {"ID_Usuario": uid, "F_Registro": date.today()})
        self.repo.insertar("T_OBJETIVOS", {
            "ID_Usuario": uid, "F_Registro": date.today(), "Tipo": "FUERZA",
            "Descripcion": "Mejorar sentadilla", "Estado": "ACTIVO"})
        self.repo.insertar("T_DISPONIBILIDAD", {
            "ID_Usuario": uid, "F_Registro": date.today(),
            "Dias_Semana": 3, "Min_Sesion": 60})
        self.assertEqual(srv.avisos(self.repo, self.repo.obtener("T_USUARIOS", uid)), [])


if __name__ == "__main__":
    unittest.main()
